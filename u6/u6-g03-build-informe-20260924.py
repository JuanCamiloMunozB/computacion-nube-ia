"""Genera el informe de proceso y hallazgos de la Unidad 6."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "u6-g03-informe-proceso-hallazgos-20260924.docx"
BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(76, 81, 88)
NAVY = "24364C"
PALE = "F3F6F8"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)


def set_cell_margin(cell, top=95, start=105, bottom=95, end=105) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def add_table(doc, headers, rows, widths, numeric_cols=()) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(header)
        set_cell_shading(cell, NAVY)
    set_repeat_header(table.rows[0])
    for row_i, values in enumerate(rows):
        cells = table.add_row().cells
        for col_i, value in enumerate(values):
            cells[col_i].text = str(value)
            if row_i % 2:
                set_cell_shading(cells[col_i], PALE)
    for row_i, row in enumerate(table.rows):
        for col_i, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margin(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                if col_i in numeric_cols:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.7 if len(headers) > 5 else 9.0)
                    run.font.color.rgb = RGBColor(255, 255, 255) if row_i == 0 else BLACK
                    if row_i == 0:
                        run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def para(doc, text: str, *, style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def heading(doc, text: str, level=1) -> None:
    doc.add_heading(text, level=level)


def bullet(doc, text: str) -> None:
    para(doc, text, style="List Bullet")


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.77)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.2)
normal.font.color.rgb = BLACK
normal.paragraph_format.line_spacing = 1.16
normal.paragraph_format.space_after = Pt(6)
for name, size, before, after in (
    ("Title", 19, 0, 10),
    ("Heading 1", 13.2, 15, 6),
    ("Heading 2", 11.1, 10, 4),
):
    style = styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = BLACK
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
styles["List Bullet"].font.name = "Arial"
styles["List Bullet"].font.size = Pt(10.2)
styles["List Bullet"].paragraph_format.space_after = Pt(3)
title_style_ppr = styles["Title"]._element.get_or_add_pPr()
title_style_border = title_style_ppr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_ppr.remove(title_style_border)

title = doc.add_paragraph(style="Title")
title.add_run("Proceso y hallazgos del monitoreo de retención en la Unidad 6")
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(13)
run = meta.add_run("Grupo 3  |  Proyecto computacionnube20263  |  24 de septiembre de 2026")
run.font.name = "Arial"
run.font.size = Pt(9.5)
run.font.color.rgb = MUTED

para(
    doc,
    "El pipeline procesó diez lotes semanales con 703 registros: 653 cumplieron el contrato y 50 quedaron en cuarentena. "
    "La calidad se deterioró con claridad el 24 y el 31 de agosto, mientras que la distribución de clientes aceptados "
    "ya mostraba cambios frente a las primeras semanas. La aplicación Streamlit reúne esas evidencias y una predicción "
    "individual probada contra la API privada. Los datos disponibles no incluyen el resultado real de las campañas ni "
    "etiquetas posteriores de churn, por lo que aún no permiten atribuir la caída de retención al modelo o a una causa de negocio concreta."
)

heading(doc, "Objetivo y alcance")
para(
    doc,
    "El trabajo respondió a las seis directrices de la Unidad 6: revisar el umbral de rechazo, inspeccionar la cuarentena, "
    "analizar los registros válidos con dos referencias, medir el cambio con límites calculados, tratar la nueva columna "
    "BancoPago y formular una respuesta defendible para el cliente. La entrega exigida por INTU es una aplicación Streamlit "
    "por grupo, con los nombres de los integrantes y capacidad de hacer predicciones individuales; no se exige un informe escrito. "
    "Este documento deja trazabilidad del proceso y sirve para preparar la sustentación."
)

heading(doc, "Datos y arquitectura implementada")
para(
    doc,
    "Trabajamos en el proyecto computacionnube20263 y usamos el prefijo u6-g03 para los recursos de la unidad. "
    "El archivo consolidado contiene diez semanas entre el 29 de junio y el 31 de agosto de 2026. Los CSV de prueba y "
    "los archivos de referencia se guardaron en Cloud Storage; el DAG de Airflow lee cada lote desde ese bucket."
)
para(
    doc,
    "El flujo primero valida el contrato de entrada. Cada fila aceptada se envía a la API de churn en Cloud Run; "
    "los rechazos se conservan con código, campos implicados y registro original en cuarentena. Después se escriben en BigQuery "
    "las predicciones, la cuarentena, el resumen de calidad por lote y el PSI por variable. La tarea final verifica si se "
    "superaron los límites de calidad o drift. Por diseño, puede terminar en FAILED después de persistir la evidencia: "
    "ese estado indica una alerta de monitoreo, no que se perdieran los resultados anteriores."
)
para(
    doc,
    "La API de Cloud Run permaneció privada. Airflow y Streamlit obtienen un ID token de la sesión autenticada de gcloud "
    "para llamar al endpoint de predicción. No se empaquetaron credenciales ni tokens. La app Streamlit también puede "
    "recalcular el análisis del CSV para explorarlo; el registro operativo de cada corrida está en BigQuery."
)

heading(doc, "Secuencia de trabajo y verificaciones")
bullet(doc, "Se revisaron el laboratorio, las directrices del cliente, el esquema de datos y el plan de implementación antes de decidir cambios.")
bullet(doc, "Se conservaron los cuatro lotes de prueba del laboratorio y se agregó el consolidado de diez semanas. El runbook documenta 20 aceptados en ok, 15 en mix, 17 en esquema y 5 en drift; las alertas de los casos problemáticos son parte del comportamiento esperado.")
bullet(doc, "Se comprobó la autenticación de gcloud y el acceso a la API privada antes de ejecutar las predicciones. El DAG guarda la evidencia en BigQuery antes de aplicar la compuerta de alerta.")
bullet(doc, "Las consultas de BigQuery confirmaron calidad por semana, variables en alerta PSI, causas de cuarentena y cobertura de BancoPago. Se contrastaron con el cálculo reproducible de los CSV locales.")
bullet(doc, "Se construyó Streamlit con seis pestañas: calidad y umbral, cuarentena, drift, BancoPago, respuesta al cliente y predicción individual. En Cloud Shell se subió solo el archivo .py; los CSV se copiaron del bucket y se reutilizó el módulo de monitoreo existente.")
bullet(doc, "La vista previa en el puerto 8501 cargó las seis pestañas. Una predicción individual de prueba devolvió HTTP 200, riesgo 42,0 % y resultado ‘Churn no probable’. Finalmente se actualizó el ZIP de entrega y se probó descomprimido: seis pestañas y ninguna excepción.")

heading(doc, "Criterios de monitoreo elegidos")
heading(doc, "Calidad y límite predictivo", 2)
para(
    doc,
    "El laboratorio partía de un umbral fijo del 30 % de registros rechazados. Usamos las primeras cuatro semanas "
    "(29 de junio a 20 de julio) como referencia estable: 237 filas, una rechazada. Cuatro semanas fue una decisión "
    "analítica explícita, no una regla impuesta por la consigna. Con una prior uniforme Beta(1,1), el código calcula "
    "el límite predictivo superior beta-binomial al 95 % para el tamaño de cada lote. Así, el número de rechazos "
    "tolerado varía con el tamaño del lote, en lugar de aplicar un porcentaje fijo a 45 y a 96 registros por igual."
)
para(
    doc,
    "El límite calibrado fue aproximadamente 2,8 % a 4,4 % según el tamaño semanal. Se mantuvo visible el umbral "
    "anterior para comparar qué decisiones habría producido, sin confundirlo con el nuevo criterio. Una alerta de calidad "
    "significa exceso de filas que no cumplen el contrato; no mide precisión del modelo ni éxito de campañas."
)
heading(doc, "Drift y referencias", 2)
para(
    doc,
    "Para las filas aceptadas se calculó PSI sobre las 12 variables del contrato. Las variables numéricas se agrupan "
    "con cortes derivados de la referencia y las categóricas por valor; se aplica suavizado para evitar divisiones por cero. "
    "El límite de cada comparación es el percentil 95 de 200 simulaciones bootstrap tomadas bajo la referencia, "
    "con el tamaño del lote actual. El método permite comparar variables de escalas diferentes y reconoce que las "
    "muestras pequeñas fluctúan más."
)
para(
    doc,
    "Se mantuvieron dos preguntas separadas. La referencia ‘primeras_semanas_produccion’ compara cada lote con los "
    "clientes aceptados en las cuatro semanas iniciales. ‘entrenamiento_telco_churn’ usa el CSV fuente completo de U4 "
    "(7043 filas) como aproximación a la población de entrenamiento; no es la partición exacta X_train. Un PSI por encima "
    "del límite detecta cambio de distribución, no pérdida demostrada de desempeño ni causalidad."
)

heading(doc, "Resultados de calidad por semana")
para(doc, "La tabla resume la evidencia consultada en BigQuery para el archivo consolidado. Los porcentajes y límites se redondean a dos decimales.")
add_table(
    doc,
    ["Fecha", "Total", "Válidos", "Cuarentena", "Rechazo", "Límite", "Alerta"],
    [
        ("29 jun", 50, 49, 1, "2,00 %", "4,00 %", "No"),
        ("06 jul", 45, 45, 0, "0,00 %", "4,44 %", "No"),
        ("13 jul", 81, 81, 0, "0,00 %", "3,70 %", "No"),
        ("20 jul", 61, 61, 0, "0,00 %", "3,28 %", "No"),
        ("27 jul", 96, 95, 1, "1,04 %", "3,13 %", "No"),
        ("03 ago", 56, 56, 0, "0,00 %", "3,57 %", "No"),
        ("10 ago", 86, 85, 1, "1,16 %", "3,49 %", "No"),
        ("17 ago", 71, 70, 1, "1,41 %", "2,82 %", "No"),
        ("24 ago", 91, 74, 17, "18,68 %", "3,30 %", "Sí"),
        ("31 ago", 66, 37, 29, "43,94 %", "3,03 %", "Sí"),
    ],
    [0.82, 0.55, 0.62, 0.91, 0.87, 0.82, 0.59],
    numeric_cols=(1, 2, 3, 4, 5, 6),
)
para(
    doc,
    "El 24 de agosto se rechazaron 17 de 91 filas y el 31 de agosto 29 de 66. Las dos semanas reúnen 46 de los "
    "50 rechazos del periodo. Con el 30 % anterior solo habría alertado el 31 de agosto; con su mitad, 15 %, "
    "habrían alertado ambas semanas; con el doble, 60 %, ninguna. El límite calibrado detecta ambas sin elegir "
    "15 % a posteriori como nuevo número arbitrario."
)

heading(doc, "Qué reveló la cuarentena")
para(
    doc,
    "Los 50 registros rechazados acumulan 55 señalamientos de campo, porque una fila puede violar más de una regla. "
    "PaymentMethod aparece 27 veces, tenure 21 y MonthlyCharges 7. En las primeras cuatro semanas solo se observó "
    "un error de MonthlyCharges. El 24 de agosto aparecieron 10 errores de PaymentMethod y 6 de tenure; "
    "el 31 de agosto, 17 y 14 respectivamente. La composición del problema cambió, no solo el volumen."
)
para(
    doc,
    "Los valores fuera del contrato en PaymentMethod fueron PayPal (7), Corporate billing (6), PSE (6), "
    "Digital wallet (5) y Credit card (manual) (3). No se deben convertir silenciosamente a una categoría antigua: "
    "primero hay que confirmar si representan medios nuevos válidos. También se debe verificar con negocio si las "
    "antigüedades mayores a 100 meses son legítimas antes de ampliar la regla. El registro crudo en cuarentena "
    "preserva evidencia, aunque dificulta consultar campos individualmente; la aplicación la extrae y agrupa para el análisis."
)

heading(doc, "Cambio entre los registros aceptados")
para(
    doc,
    "Superar la validación no implica parecerse a la población inicial. Frente a las primeras semanas, el primer "
    "PSI fuera de límite para tenure apareció el 10 de agosto. El 17 de agosto, tenure y MonthlyCharges quedaron "
    "simultáneamente en alerta y siguieron así hasta el último lote. La comparación con el CSV fuente de U4 "
    "ya alertaba en la primera semana para dos variables, mientras que la referencia de producción no alertaba "
    "entonces. En todo el periodo hubo 35 alertas variable-lote frente al proxy de U4 y 8 frente a producción."
)
add_table(
    doc,
    ["Fecha", "PSI tenure", "Límite", "PSI cargo", "Límite", "Lectura"],
    [
        ("10 ago", "0,409", "0,344", "0,131", "0,357", "Tenure"),
        ("17 ago", "1,179", "0,495", "0,894", "0,442", "Ambas"),
        ("24 ago", "2,164", "0,433", "1,333", "0,437", "Ambas"),
        ("31 ago", "3,265", "0,798", "2,448", "0,861", "Ambas"),
    ],
    [0.83, 1.10, 0.82, 1.12, 0.82, 1.62],
    numeric_cols=(1, 2, 3, 4),
)
para(
    doc,
    "Como medida de dirección y magnitud, la mediana de tenure entre los 236 clientes aceptados de la referencia "
    "inicial fue 32 meses y entre los 37 aceptados del 31 de agosto fue 70 meses. La mediana de MonthlyCharges "
    "pasó de 65,93 a 116,92 en las unidades del CSV. Estos perfiles describen a los registros válidos; "
    "no corrigen por posibles cambios en la selección comercial de clientes."
)

heading(doc, "Decisión sobre BancoPago")
bank_paragraph = para(
    doc,
    "BancoPago no aparecía en las primeras cuatro semanas. Desde el 27 de julio hubo 466 filas y el campo "
    "estuvo informado en 444, una cobertura de 95,3 %. Seis valores informados quedaron como ‘Desconocido’ "
    "tras normalizar variantes conocidas de los bancos. Mantuvimos el valor original y la categoría normalizada "
    "para seguimiento operativo. No imputamos las semanas anteriores ni incorporamos BancoPago al modelo existente, "
    "porque no forma parte de su contrato de predicción."
)
bank_paragraph.paragraph_format.keep_together = True

heading(doc, "Interpretación para el cliente y recomendación")
para(
    doc,
    "Los hechos permiten afirmar que llegó una población de clientes distinta, con mayor antigüedad y cargo mensual "
    "entre los aceptados, al mismo tiempo que aumentaron los pagos fuera de contrato y los rechazos. Una hipótesis "
    "plausible es que cambió la selección de clientes en el CRM, el catálogo de pagos o ambos. La coincidencia temporal "
    "no demuestra cuál mecanismo hizo caer el desempeño de las campañas."
)
para(
    doc,
    "Recomendamos continuar temporalmente el scoring solo para filas que cumplen el contrato, con monitoreo y "
    "revisión humana de las alertas. Los rechazos deben corregirse o recapturarse, no forzarse a categorías antiguas. "
    "Al cliente hay que solicitar la regla semanal de selección, fechas de cambios en CRM y pagos, acción de campaña "
    "por cliente, fecha de contacto, resultado de retención y etiqueta posterior de churn. Con datos recientes "
    "etiquetados se podrá medir desempeño, calibración y costo de negocio del modelo actual y comparar un eventual "
    "reentrenamiento antes de desplegarlo."
)
para(
    doc,
    "Sin esas etiquetas y resultados no puede concluirse que el modelo perdió precisión, que el drift causó "
    "la caída de retención o que la campaña dejó de funcionar por una sola variable. Tampoco debe interpretarse "
    "la probabilidad individual que devuelve la API como churn observado."
)

heading(doc, "Entregables y estado al cierre de esta revisión")
bullet(doc, "El DAG, los módulos de monitoreo y la aplicación Streamlit están en la carpeta u6. La evidencia operativa está en las tablas de resultados, cuarentena, monitoreo y drift del dataset u6_g03_data_20260924 de BigQuery.")
bullet(doc, "El ZIP u6-g03-streamlit-entrega-20260924.zip incluye la versión actualizada de la app, módulos, CSV de referencia, requisitos e instrucciones. La prueba tras descomprimirlo mostró seis pestañas y cero excepciones.")
bullet(doc, "La app respondió a una predicción individual desde la sesión de Cloud Shell del grupo. Todavía debe verificarse que la identidad o el entorno de la persona evaluadora pueda invocar la API privada; la URL de vista previa de Cloud Shell es temporal.")
bullet(doc, "No consta aún el envío por INTU ni la coordinación por correo de la sustentación. La consigna exige mantener los servicios disponibles hasta el cierre indicado para el viernes 25 de septiembre.")

heading(doc, "Fuentes y rutas de verificación")
para(
    doc,
    "Consigna y laboratorio: Monitoreo_pipeline_retencion.md y UNIDAD6_Lab_Paso_a_Paso (1).md. "
    "Plan y ejecución: docs/superpowers/plans/2026-09-24-u6-g03-entrega-streamlit.md y "
    "u6/u6-g03-ejecucion-20260924.md. Código: u6/dag_pipeline_churn.py, u6/streamlit_churn_u6.py, "
    "monitoring.py y delivery.py. Datos: u6-g03-lotes-retencion-20260924.csv y "
    "u6-g03-telco-churn-20260924.csv."
)
para(
    doc,
    "Evidencia operativa en el dataset computacionnube20263.u6_g03_data_20260924: tablas "
    "u6_g03_resultados_20260924, u6_g03_cuarentena_20260924, u6_g03_monitoreo_lotes_20260924 y "
    "u6_g03_drift_lotes_20260924. Las cifras semanales se contrastaron con las consultas de BigQuery "
    "y el recálculo local del CSV."
)

doc.core_properties.title = "Proceso y hallazgos del monitoreo de retención en la Unidad 6"
doc.core_properties.subject = "Unidad 6 Grupo 3 monitoreo de churn"
doc.core_properties.keywords = "Unidad 6, churn, monitoreo, Airflow, BigQuery, Streamlit"
doc.save(OUTPUT)
print(OUTPUT)
